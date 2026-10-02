"""Data workflow on a project: versioning, row actions, decisions, application to the model, provenance, approval,
historical reproducibility, CLI and UI. SYNTHETIC TEST DATA only."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest

pytest.importorskip("scipy")
pytest.importorskip("openpyxl")

from typer.testing import CliRunner  # noqa: E402

from simforge.cli import app as cli  # noqa: E402
from simforge.data.dataset import Decision  # noqa: E402
from simforge.data.service import DataDecisionRequired  # noqa: E402
from simforge.data.store import DatasetError  # noqa: E402
from simforge.domain.io import load_model  # noqa: E402
from simforge.domain.values import ValueStatus  # noqa: E402
from simforge.experiments.runner import run_simulation  # noqa: E402
from simforge.services.app import ApprovalRequired, SimForgeApp, data_linked_paths  # noqa: E402

ROOT = Path(__file__).parents[1]
SYN = ROOT / "examples" / "data" / "synthetic"
E2E_MODEL = ROOT / "examples" / "data" / "e2e_selective_per_circuit.yaml"
XLSX = SYN / "tiempos_montaje.xlsx"
TARGET = "nodes.assembly.params.process_time"
runner = CliRunner()


@pytest.fixture
def env(tmp_path):
    sf = SimForgeApp(workspace=tmp_path / "ws", library_dir=tmp_path / "lib", provider=None)
    p = sf.create_project("datos")
    sf.save_model(p, load_model(E2E_MODEL), "base")
    sf.approve_model(p, by="ana")
    return sf, p, sf.data(p)


def _import_e2e(ds, **kw):
    args = dict(sheet="Hoja1", unit="s", timestamp="Fecha", synthetic=True, by="ana")
    args.update(kw)
    return ds.import_file(XLSX, "montaje", "Tiempo", "PROCESSING_TIME", "PER_CIRCUIT", **args)


# ------------------------------------------------------------------------------------------- versioning
def test_dataset_is_immutable_hashed_and_versioned(env, tmp_path):
    sf, p, ds = env
    r = _import_e2e(ds)
    m = r.meta
    assert m.dataset_id == "montaje@v1" and m.synthetic and m.n_rows == 87 and m.counts["INVALID"] == 1
    d = p.root / "datasets" / "montaje" / "v0001"
    assert (d / "source" / "tiempos_montaje.xlsx").read_bytes() == XLSX.read_bytes()  # byte-identical copy
    import stat
    for f in ("dataset.json", "observations.json", "source/tiempos_montaje.xlsx"):
        assert not (d / f).stat().st_mode & (stat.S_IWUSR | stat.S_IWGRP | stat.S_IWOTH)  # stored read-only
    # same file + same settings -> duplicate, nothing new
    again = _import_e2e(ds)
    assert again.duplicate_of == "montaje@v1" and len(ds.list()) == 1
    # same file, other interpretation -> new version, noted
    other = _import_e2e(ds, unit="min")
    assert other.meta.dataset_id == "montaje@v2" and other.same_file_other_settings == ["montaje@v1"]
    # modified file -> new version, v1 untouched
    mod = tmp_path / "tiempos_montaje.xlsx"
    from openpyxl import load_workbook
    wb = load_workbook(XLSX)
    wb["Hoja1"]["C2"] = 999.0
    wb.save(mod)
    v3 = ds.import_file(mod, "montaje", "Tiempo", "PROCESSING_TIME", "PER_CIRCUIT", sheet="Hoja1", unit="s", synthetic=True)
    assert v3.meta.dataset_id == "montaje@v3" and v3.meta.content_hash != m.content_hash
    assert ds.store.meta("montaje@v1").content_hash == m.content_hash
    forced = _import_e2e(ds, force_new_version=True)
    assert forced.meta.dataset_id == "montaje@v4"


def test_synthetic_file_name_cannot_be_imported_as_measured(env):
    _, _, ds = env
    m = ds.import_file(SYN / "SYNTHETIC_E_outliers.csv", "e", "tiempo_s", "PROCESSING_TIME", "PER_UNIT", unit="s").meta
    assert m.synthetic and m.label.startswith("[SYNTHETIC TEST DATA]")


# -------------------------------------------------------------------------------------------- row actions
def test_outlier_keep_exclude_restore_are_logged_and_reversible(env):
    _, p, ds = env
    did = ds.import_file(SYN / "SYNTHETIC_E_outliers.csv", "e", "tiempo_s", "PROCESSING_TIME", "PER_UNIT", unit="s").meta.dataset_id
    prof = ds.profile(did)
    assert {10, 40, 70} <= {c["index"] for c in prof["outliers"]["candidates"]}
    assert prof["status"] == "REQUIRES_ENGINEER_REVIEW" and prof["stats"]["n"] == 80
    ds.act_rows(did, "KEEP", [10, 40], "atascos reales", by="ana", method="IQR")
    ds.act_rows(did, "EXCLUDE_FROM_FIT", [70], "5 s: error de registro (lector duplicado)", by="ana", method="MAD")
    st = ds.state(did)
    assert st.excluded == [70] and 70 not in st.used and len(st.used) == 79
    assert len(ds.store.observations(did)) == 80  # original data never removed
    h1 = st.state_hash
    ds.act_rows(did, "RESTORED", [70], "revisado: el 5 s era real", by="luis")
    st = ds.state(did)
    assert st.excluded == [] and len(st.used) == 80 and st.state_hash != h1
    assert [e.action for e in st.row_log] == ["KEEP", "EXCLUDE_FROM_FIT", "RESTORED"]
    assert all(e.by and e.reason and e.at for e in st.row_log)
    with pytest.raises(DatasetError):
        ds.act_rows(did, "RESTORED", [70], "nothing to restore", by="ana")
    with pytest.raises(Exception):
        ds.act_rows(did, "EXCLUDE_FROM_FIT", [1], "", by="ana")  # a reason is mandatory


def test_review_rows_are_not_used_until_accepted_and_invalid_never(env):
    _, _, ds = env
    did = ds.import_file(SYN / "SYNTHETIC_D_missing.csv", "d", "tiempo_s", "PROCESSING_TIME", "PER_UNIT", unit="s").meta.dataset_id
    st = ds.state(did)
    assert st.pending_review == [33] and 33 not in st.used and len(st.invalid) == 6
    with pytest.raises(DatasetError):
        ds.act_rows(did, "RESTORED", [25], "negative", by="ana")  # import-INVALID rows cannot be brought back
    with pytest.raises(DatasetError):
        ds.act_rows(did, "ACCEPT_REVIEWED", [25], "not pending", by="ana")
    ds.act_rows(did, "ACCEPT_REVIEWED", [33], "0 s real: pieza ya montada", by="ana")
    assert 33 in ds.state(did).used


# ---------------------------------------------------------------------------------------------- groups
def test_groups_are_never_mixed_silently(env):
    _, _, ds = env
    did = ds.import_file(SYN / "SYNTHETIC_H_operator_product.csv", "h", "tiempo_s", "PROCESSING_TIME", "PER_UNIT", unit="s",
                         groups={"operator": "operario", "product": "producto"}).meta.dataset_id
    prof = ds.profile(did)
    assert "selection_error" in prof and prof["groups"]["operator"]["flag"] == "GROUPS_DIFFER"
    with pytest.raises(DataDecisionRequired, match="mezclan grupos"):
        ds.fit(did, by="ana")
    a = ds.fit(did, by="ana", group={"operator": "A", "product": "X"})
    assert a["n"] == 40 and a["group"] == {"operator": "A", "product": "X"}
    pooled = ds.fit(did, by="ana", pooled=True)
    assert pooled["n"] == 160 and pooled["pooled"] is True


# -------------------------------------------------------------------------------------- decide and apply
def test_decisions_deterministic_empirical_fitted(env):
    _, _, ds = env
    did = _import_e2e(ds).meta.dataset_id
    with pytest.raises(DataDecisionRequired):
        ds.decide(did, "USE_DETERMINISTIC", by="ana")  # MEAN, MEDIAN or ENGINEER_VALUE must be stated
    mean = ds.decide(did, "USE_DETERMINISTIC", by="ana", derivation="MEAN")
    med = ds.decide(did, "USE_DETERMINISTIC", by="ana", derivation="MEDIAN")
    eng = ds.decide(did, "USE_DETERMINISTIC", by="ana", derivation="ENGINEER_VALUE", engineer_value=30)
    assert mean.distribution["value"] != med.distribution["value"] and eng.distribution == {"dist": "constant", "value": 30.0, "unit": "s"}
    emp = ds.decide(did, "USE_EMPIRICAL", by="ana")
    assert emp.distribution["dist"] == "empirical" and len(emp.distribution["values"]) == 86 and emp.n_used == 86
    fit = ds.fit(did, by="ana")
    with pytest.raises(DataDecisionRequired, match="Avisos de plausibilidad"):
        ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="exponential")
    with pytest.raises(DataDecisionRequired, match="truncamiento"):
        ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="normal")
    with pytest.raises(DataDecisionRequired, match="bound_type"):
        ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="normal",
                  truncation={"lower": 0, "reason": "tiempos no negativos"}, accept_warnings=["MASS_BELOW_OBSERVED_MIN"])
    nt = ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="normal",
                   truncation={"lower": 0, "reason": "tiempos no negativos", "bound_type": "MODELLING_BOUND"},
                   accept_warnings=["MASS_BELOW_OBSERVED_MIN"])
    tr = nt.distribution["truncation"]
    assert tr["lower"] == 0 and tr["method"] == "FIT_THEN_TRUNCATE" and tr["provenance"]["source"] == "engineer:ana"
    assert any(w.startswith("FIT_THEN_TRUNCATE") for w in nt.warnings)
    ln = ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="lognormal")
    assert ln.decision is Decision.USE_FITTED and ln.decision_id == "dec_006"
    # data changed after the fit -> the fit is stale
    ds.act_rows(did, "EXCLUDE_FROM_FIT", [20], "prueba", by="ana")
    with pytest.raises(DataDecisionRequired, match="vuelve a ajustar"):
        ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="lognormal")
    ds.decide(did, "KEEP_WITHOUT_APPLYING", by="ana", reason="faltan datos del turno de tarde")
    assert ds.profile(did)["status"] == "REQUIRES_ENGINEER_REVIEW"  # outlier 61 still not reviewed


def test_apply_creates_version_with_provenance_and_invalidates_approval(env):
    sf, p, ds = env
    did = _import_e2e(ds, synthetic=False).meta.dataset_id  # pretend measured, to check MEASURED provenance
    assert p.current_model().is_approved
    fit = ds.fit(did, by="ana")
    dec = ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="gamma")
    with pytest.raises(DataDecisionRequired, match="PROPUESTA: sum_iid"):  # proposed, never inferred
        ds.apply(did, dec.decision_id, TARGET, "PER_CIRCUIT", by="ana")
    with pytest.raises(DataDecisionRequired, match="Base del dato"):
        ds.apply(did, dec.decision_id, TARGET, "PER_RACK", by="ana", aggregation="sum_iid")
    before = p.meta.current_version
    r = ds.apply(did, dec.decision_id, TARGET, "PER_CIRCUIT", by="ana", aggregation="sum_iid")
    assert r["version"] == before + 1 and r["was_approved"] is True
    assert r["work_units_aggregation_change"] == (None, "sum_iid")
    assert p.current_model().nodes[1].params["work_units_aggregation"] == "sum_iid"
    assert ds.state(did).applications[-1].work_units_aggregation == "sum_iid"
    m = p.current_model()
    assert not m.is_approved and "invalidated" in m.approval.note
    pt = m.nodes[1].params["process_time"]
    assert pt["dist"] == "gamma" and pt["unit"] == "s"
    prov = pt["provenance"]
    assert prov["status"] == ValueStatus.MEASURED.value
    link = prov["data"]
    assert link["dataset_id"] == did and link["fit_id"] == fit["fit_id"] and link["basis"] == "PER_CIRCUIT"
    assert link["content_hash"] == ds.store.meta(did).content_hash and link["decision"] == "FITTED" and link["n_used"] == 86
    assert data_linked_paths(m) == [TARGET]
    with pytest.raises(ApprovalRequired, match="derivados de datos importados"):
        sf.run_simulation(p)
    sf.approve_model(p, by="ana")
    res = sf.run_simulation(p, replications=2)
    assert res.kpis.mean("units_completed") > 0
    hist = [h["action"] for h in p.history()]
    assert "data_apply" in hist and "data_decision" in hist and "data_import" in hist


def test_basis_conversion_is_explicit_and_calculated(env):
    sf, p, ds = env
    model = p.current_model()
    from simforge.domain.paths import set_value
    sf.save_model(p, set_value(model, "nodes.assembly.params.work_units", 1), "per rack")
    did = _import_e2e(ds, synthetic=False).meta.dataset_id
    mean = ds.decide(did, "USE_DETERMINISTIC", by="ana", derivation="MEAN")
    with pytest.raises(DataDecisionRequired, match="conversión explícita"):
        ds.apply(did, mean.decision_id, TARGET, "PER_RACK", by="ana")
    conv = {"factor": 4, "formula": "t_rack = 4 · t_circuito", "inputs": {"circuits_per_rack": 4}}
    r = ds.apply(did, mean.decision_id, TARGET, "PER_RACK", by="ana", conversion=conv)
    pt = p.current_model().nodes[1].params["process_time"]
    assert pt["value"] == pytest.approx(4 * mean.distribution["value"]) and pt["provenance"]["status"] == "calculated"
    assert pt["provenance"]["data"]["conversion"]["to_basis"] == "PER_RACK" and "CONVERSION" in r["warnings"][0]
    emp = ds.decide(did, "USE_EMPIRICAL", by="ana")
    with pytest.raises(DataDecisionRequired, match="scale_sample"):
        ds.apply(did, emp.decision_id, TARGET, "PER_RACK", by="ana", conversion=conv)
    r = ds.apply(did, emp.decision_id, TARGET, "PER_RACK", by="ana", conversion={**conv, "aggregation": "scale_sample"})
    pt = p.current_model().nodes[1].params["process_time"]
    assert pt["dist"] == "empirical" and pt["provenance"]["data"]["conversion"]["aggregation"] == "scale_sample"


def test_apply_refuses_unsupported_targets_and_quantities(env):
    _, p, ds = env
    did = _import_e2e(ds).meta.dataset_id
    dec = ds.decide(did, "USE_DETERMINISTIC", by="ana", derivation="MEDIAN")
    with pytest.raises(DatasetError, match="solo puede aplicarse"):
        ds.apply(did, dec.decision_id, "nodes.assembly.params.capacity", "PER_CIRCUIT", by="ana")
    with pytest.raises(DatasetError, match="no admite"):
        ds.apply(did, dec.decision_id, "nodes.selective_in.params.process_time", "PER_CIRCUIT", by="ana")
    with pytest.raises(DatasetError, match="averías"):
        r = ds.import_file(SYN / "SYNTHETIC_A_constant.csv", "rep", "tiempo_s", "REPAIR_TIME", "PER_OPERATION", unit="s")
        d2 = ds.decide(r.meta.dataset_id, "USE_DETERMINISTIC", by="ana", derivation="MEAN")
        ds.apply(r.meta.dataset_id, d2.decision_id, "nodes.assembly.params.failures.mttr", "PER_OPERATION", by="ana")
    s = ds.import_file(SYN / "SYNTHETIC_A_constant.csv", "setup", "tiempo_s", "SETUP_TIME", "PER_BATCH", unit="s")
    d3 = ds.decide(s.meta.dataset_id, "USE_DETERMINISTIC", by="ana", derivation="MEAN")
    with pytest.raises(DatasetError, match="SIMULATION_USE_NOT_YET_SUPPORTED"):
        ds.apply(s.meta.dataset_id, d3.decision_id, "nodes.assembly.params.process_time", "PER_BATCH", by="ana")
    u = ds.import_file(SYN / "SYNTHETIC_A_constant.csv", "unk", "tiempo_s", "PROCESSING_TIME", "UNKNOWN", unit="s")
    d4 = ds.decide(u.meta.dataset_id, "USE_DETERMINISTIC", by="ana", derivation="MEAN")
    with pytest.raises(DataDecisionRequired, match="UNKNOWN"):
        ds.apply(u.meta.dataset_id, d4.decision_id, TARGET, "PER_CIRCUIT", by="ana")
    rej = ds.decide(did, "REJECT", by="ana", reason="estudio mal hecho")
    with pytest.raises(DatasetError, match="no se aplica"):
        ds.apply(did, rej.decision_id, TARGET, "PER_CIRCUIT", by="ana")
    assert ds.profile(did)["status"] == "REJECTED"


def test_arrival_interval_and_failures_targets(env):
    sf, p, ds = env
    from simforge.domain.paths import set_value
    m = p.current_model()
    m = set_value(m, "nodes.src.params.arrival", "interarrival")
    m = set_value(m, "nodes.src.params.interarrival", {"dist": "constant", "value": 100})
    m = set_value(m, "nodes.selective.params.failures", {"mtbf": {"dist": "constant", "value": 2, "unit": "h"},
                                                         "mttr": {"dist": "constant", "value": 5, "unit": "min"}})
    sf.save_model(p, m, "arrivals + failures")
    c = ds.import_file(SYN / "SYNTHETIC_C_arrivals.csv", "arr", "intervalo_s", "ARRIVAL_INTERVAL", "PER_UNIT", unit="s",
                       timestamp="llegada")
    fit = ds.fit(c.meta.dataset_id, by="ana")
    assert candidate_rank(fit, "exponential") <= 3
    d = ds.decide(c.meta.dataset_id, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="exponential",
                  accept_warnings=["TAIL_EXTRAPOLATION", "EXTREME_TAIL", "MASS_BELOW_OBSERVED_MIN"])
    ds.apply(c.meta.dataset_id, d.decision_id, "nodes.src.params.interarrival", "PER_UNIT", by="ana")
    r = ds.import_file(SYN / "SYNTHETIC_B_skewed.csv", "rep", "tiempo (s)", "REPAIR_TIME", "PER_OPERATION", unit="s")
    d2 = ds.decide(r.meta.dataset_id, "USE_EMPIRICAL", by="ana")
    ds.apply(r.meta.dataset_id, d2.decision_id, "nodes.selective.params.failures.mttr", "PER_OPERATION", by="ana")
    m = p.current_model()
    assert m.nodes[0].params["interarrival"]["dist"] == "exponential"
    assert m.nodes[3].params["failures"]["mttr"]["dist"] == "empirical"
    sf.approve_model(p, by="ana")
    assert sf.run_simulation(p, replications=2).kpis.mean("units_completed") > 0


def candidate_rank(fit, fam):
    return next(c["rank"] for c in fit["candidates"] if c["family"] == fam)


def test_distance_only_deterministic(env):
    _, p, ds = env
    d = ds.import_file(SYN / "SYNTHETIC_A_constant.csv", "dist", "tiempo_s", "DISTANCE", "PER_TRIP", unit="cm")
    assert d.meta.analysis_unit == "cm" and d.meta.normalized_unit == "m"
    emp = ds.decide(d.meta.dataset_id, "USE_EMPIRICAL", by="ana")
    with pytest.raises(DatasetError, match="no admite|nodes"):  # e2e model has no transport node
        ds.apply(d.meta.dataset_id, emp.decision_id, "nodes.assembly.params.distance", "PER_TRIP", by="ana")


# --------------------------------------------------------------------------------- reproducibility
def test_historical_run_is_reproducible_after_source_changes(env, tmp_path):
    sf, p, ds = env
    did = _import_e2e(ds).meta.dataset_id
    emp = ds.decide(did, "USE_EMPIRICAL", by="ana")
    v = ds.apply(did, emp.decision_id, TARGET, "PER_CIRCUIT", by="ana", aggregation="sum_iid")["version"]
    sf.approve_model(p, by="ana")
    run = sf.run_simulation(p, replications=3)
    # the Excel changes later and is re-imported: a NEW dataset version; the stored model version is untouched
    mod = tmp_path / "tiempos_montaje.xlsx"
    shutil.copy(XLSX, mod)
    from openpyxl import load_workbook
    wb = load_workbook(mod)
    wb["Hoja1"]["C3"] = 500.0
    wb.save(mod)
    assert ds.import_file(mod, "montaje", "Tiempo", "PROCESSING_TIME", "PER_CIRCUIT", sheet="Hoja1", unit="s",
                          synthetic=True).meta.dataset_id == "montaje@v2"
    old = p.load_version(v + 1)  # the approved version that ran
    again = run_simulation(old, sf.registry, replications=3)  # no cache: recomputed from version + seeds + engine
    assert again.model_hash == run.model_hash and again.seeds == run.seeds and again.engine_version == run.engine_version
    assert again.kpis.mean("units_completed") == run.kpis.mean("units_completed")
    assert again.kpis.mean("avg_lead_time_s") == run.kpis.mean("avg_lead_time_s")
    t = ds.trace_parameter(old, TARGET)
    assert t["data_link"]["dataset_id"] == "montaje@v1" and t["dataset_hash_matches"] is True


def test_data_derived_parameters_change_the_model_hash_but_not_example_hashes(env):
    _, p, ds = env
    did = _import_e2e(ds).meta.dataset_id
    h0 = p.current_model().content_hash()
    d = ds.decide(did, "USE_DETERMINISTIC", by="ana", derivation="ENGINEER_VALUE", engineer_value=30)
    ds.apply(did, d.decision_id, TARGET, "PER_CIRCUIT", by="ana")
    m = p.current_model()
    assert m.content_hash() != h0  # same value (30 s) but now traceable to data -> a different, re-approvable model
    assert m.nodes[1].params["process_time"]["provenance"]["status"] == "assumed"  # synthetic data is never MEASURED


# ------------------------------------------------------------------------------------------------ CLI
def test_cli_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setenv("SIMFORGE_WORKSPACE", str(tmp_path / "ws"))
    monkeypatch.setenv("SIMFORGE_LIBRARY", str(tmp_path / "lib"))

    def ok(*args):
        r = runner.invoke(cli, list(map(str, args)))
        assert r.exit_code == 0, r.output
        return r.output
    ok("project", "new", "E2E", "--model", E2E_MODEL)
    ok("project", "approve", "e2e", "--by", "ana")
    assert "Sheets: ['LEEME', 'Hoja1']" in ok("data", "preview", XLSX)
    out = ok("data", "import", "e2e", XLSX, "--name", "montaje", "--sheet", "Hoja1", "--column", "Tiempo", "--unit", "s",
             "--timestamp", "Fecha", "--quantity", "PROCESSING_TIME", "--basis", "PER_CIRCUIT", "--synthetic", "--by", "ana")
    assert "[SYNTHETIC TEST DATA] montaje@v1" in out
    out = ok("data", "inspect", "e2e", "montaje@v1")
    assert "BASIS:        PER_CIRCUIT" in out and "POTENTIAL OUTLIERS: 2" in out and "REQUIRES_ENGINEER_REVIEW" in out
    ok("data", "rows", "e2e", "montaje@v1", "KEEP", "20", "61", "--reason", "ciclos largos reales", "--by", "ana")
    out = ok("data", "fit", "e2e", "montaje@v1", "--by", "ana")
    assert "SUGGESTED CANDIDATE" in out and "REQUIRES ENGINEER ACCEPTANCE" in out and "Nada se ha aplicado" in out
    fit_id = out.split("(fit ")[1].split(",")[0]
    ok("data", "decide", "e2e", "montaje@v1", "USE_FITTED", "--fit", fit_id, "--candidate", "lognormal", "--by", "ana")
    r = runner.invoke(cli, ["data", "apply", "e2e", "montaje@v1", "dec_001", "--target", TARGET, "--target-basis", "PER_CIRCUIT",
                            "--by", "ana"])
    assert r.exit_code == 1 and "PROPUESTA: sum_iid" in r.output
    out = ok("data", "apply", "e2e", "montaje@v1", "dec_001", "--target", TARGET, "--target-basis", "PER_CIRCUIT", "--by", "ana",
             "--aggregation", "sum_iid")
    assert "INVALIDATED" in out
    r = runner.invoke(cli, ["project", "run", "e2e"])
    assert r.exit_code == 1 and "derivados de datos" in r.output
    ok("project", "approve", "e2e", "--by", "ana")
    out = ok("project", "run", "e2e", "--reps", "2")
    assert "engine 0.6.0" in out and "seeds [12345, 12346]" in out
    assert '"fit_id": "' + fit_id in ok("data", "trace", "e2e", TARGET)
    assert "APPLIED" in ok("data", "inspect", "e2e", "montaje@v1")
    assert "montaje@v1" in ok("data", "list", "e2e")
    assert "SIMULATION_USE_NOT_YET_SUPPORTED" in ok("data", "quantities")
    r = runner.invoke(cli, ["data", "import", "e2e", str(SYN / "SYNTHETIC_D_missing.csv"), "--name", "x", "--column", "tiempo_s",
                            "--quantity", "PROCESSING_TIME", "--basis", "PER_UNIT"])
    assert r.exit_code == 1 and "unidad" in r.output  # unit is mandatory: no traceback, a message


# ------------------------------------------------------------------------------------------------- UI
def test_ui_data_tab_renders_profile_and_fits(tmp_path, monkeypatch):
    pytest.importorskip("streamlit")
    from streamlit.testing.v1 import AppTest

    ws = tmp_path / "ws"
    sf = SimForgeApp(workspace=ws, library_dir=tmp_path / "lib", provider=None)
    p = sf.create_project("UI data")
    sf.save_model(p, load_model(E2E_MODEL), "base")
    _import_e2e(sf.data(p))
    monkeypatch.setenv("SIMFORGE_WORKSPACE", str(ws))
    monkeypatch.setenv("SIMFORGE_LIBRARY", str(tmp_path / "lib"))
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.chdir(tmp_path)
    at = AppTest.from_file(str(ROOT / "src" / "simforge" / "ui" / "app.py"), default_timeout=120)
    at.run()
    assert not at.exception
    assert any("DATASET:      montaje@v1" in c.value for c in at.code)
    next(t for t in at.text_input if t.key == "eng").input("ana").run()
    next(b for b in at.button if b.label == "Ajustar").click().run()
    assert not at.exception
    assert any("STATISTICAL RANKING" in c.value for c in at.code)


def test_experiment_with_data_derived_distribution_uses_recorded_seeds(env):
    """Experiments keep working with distributed parameters: every scenario uses the same seeds per replication
    (common random numbers through per-(seed, node, purpose) streams)."""
    sf, p, ds = env
    did = _import_e2e(ds).meta.dataset_id
    fit = ds.fit(did, by="ana")
    d = ds.decide(did, "USE_FITTED", by="ana", fit_id=fit["fit_id"], candidate="gamma")
    ds.apply(did, d.decision_id, TARGET, "PER_CIRCUIT", by="ana", aggregation="sum_iid")
    sf.approve_model(p, by="ana")
    from simforge.domain.isms import ExperimentSpec, Factor
    spec = ExperimentSpec(name="racks", factors=[Factor(path="resources.racks.quantity", values=[2, 4])], replications=3)
    a = sf.run_experiment(p, spec)
    b = sf.run_experiment(p, spec)
    assert [s.result.seeds for s in a.scenarios] == [[12345, 12346, 12347]] * 2
    assert a.table(["units_completed"]) == b.table(["units_completed"])
    assert a.scenarios[0].result.kpis.stats["units_completed"].half_width > 0  # the data-derived variability is there
